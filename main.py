import argparse
import csv
import hashlib
import json
import re
import sqlite3
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

import feedparser
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

BASE_DIR = Path(__file__).resolve().parent
DB_FILE = BASE_DIR / "news.db"
JSON_FILE = BASE_DIR / "news.json"

FEEDS = {
    "BBC News": "https://feeds.bbci.co.uk/news/rss.xml",
    "The Guardian": "https://www.theguardian.com/world/rss",
    "NPR": "https://feeds.npr.org/1001/rss.xml",
}

FIELDS = ["title", "source", "published", "description", "url", "article_id"]


def connect_db():
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    conn.execute("""
        CREATE TABLE IF NOT EXISTS articles (
            article_id TEXT PRIMARY KEY,
            title TEXT NOT NULL,
            source TEXT NOT NULL,
            published TEXT,
            description TEXT,
            url TEXT UNIQUE
        )
    """)
    conn.commit()
    return conn


def clean_url(url):
    if not url:
        return ""
    parts = urlsplit(url)
    return urlunsplit((
        parts.scheme, parts.netloc.lower(), parts.path.rstrip("/"), "", ""
    ))


def clean_text(text):
    text = re.sub(r"<[^>]+>", " ", text or "")
    return re.sub(r"\s+", " ", text).strip()


def publication_date(entry):
    for field in ("published_parsed", "updated_parsed"):
        value = entry.get(field)
        if value:
            try:
                return datetime(
                    *value[:6], tzinfo=timezone.utc
                ).isoformat()
            except (TypeError, ValueError, OverflowError):
                pass
    return ""


def fetch_news():
    conn = connect_db()
    added = 0

    try:
        for source, feed_url in FEEDS.items():
            print(f"Fetching {source}...")
            try:
                import feedparser
                feed = feedparser.parse(feed_url)

                for entry in feed.entries:
                    title = clean_text(entry.get("title", ""))
                    if not title:
                        continue

                    url = clean_url(entry.get("link", ""))
                    identity = url or title.lower()
                    article_id = hashlib.sha256(
                        identity.encode("utf-8")
                    ).hexdigest()

                    cursor = conn.execute("""
                        INSERT OR IGNORE INTO articles
                        (article_id, title, source, published, description, url)
                        VALUES (?, ?, ?, ?, ?, ?)
                    """, (
                        article_id, title, source,
                        publication_date(entry),
                        clean_text(entry.get("summary", "")),
                        url or None,
                    ))
                    added += max(cursor.rowcount, 0)

                conn.commit()
                print(f"Checked {len(feed.entries)} feed entries.")
            except Exception as error:
                print(f"Could not fetch {source}: {error}")

        rows = conn.execute(
            "SELECT * FROM articles ORDER BY published DESC"
        ).fetchall()
        articles = [dict(row) for row in rows]

        with open(JSON_FILE, "w", encoding="utf-8") as file:
            json.dump(articles, file, indent=2, ensure_ascii=False)

        print(f"New articles saved: {added}")
        print(f"Total articles: {len(articles)}")
        print("Saved in SQLite and JSON.")
    finally:
        conn.close()


def add_filters(parser):
    parser.add_argument("--source")
    parser.add_argument("--keyword")
    parser.add_argument("--from-date", dest="from_date")
    parser.add_argument("--to-date", dest="to_date")


def get_articles(args):
    for name in ("from_date", "to_date"):
        value = getattr(args, name, None)
        if value:
            try:
                date.fromisoformat(value)
            except ValueError:
                raise ValueError(f"{name} must use YYYY-MM-DD")

    if args.from_date and args.to_date and args.from_date > args.to_date:
        raise ValueError("Start date must not be after end date.")

    conn = connect_db()
    try:
        sql = "SELECT * FROM articles WHERE 1=1"
        params = []

        if args.source:
            sql += " AND source LIKE ?"
            params.append(f"%{args.source}%")

        if args.keyword:
            sql += " AND (title LIKE ? OR description LIKE ?)"
            term = f"%{args.keyword}%"
            params.extend([term, term])

        if args.from_date:
            sql += " AND substr(published,1,10) >= ?"
            params.append(args.from_date)

        if args.to_date:
            sql += " AND substr(published,1,10) <= ?"
            params.append(args.to_date)

        sql += " ORDER BY published DESC"
        return [dict(row) for row in conn.execute(sql, params).fetchall()]
    finally:
        conn.close()


def list_news(args):
    articles = get_articles(args)
    if not articles:
        print("No articles found. Run 'python main.py fetch' first.")
        return

    for number, article in enumerate(articles, 1):
        print(f"\n{number}. {article['title']}")
        print(f"Source: {article['source']}")
        print(f"Published: {article['published'] or 'Unknown'}")
        print(f"URL: {article['url']}")
        print(f"Summary: {article['description'][:200]}")


def export_news(args):
    articles = get_articles(args)
    if not articles:
        print("No articles to export.")
        return

    output = Path(args.output)
    if not output.is_absolute():
        output = BASE_DIR / output

    if output.suffix.lower() == ".csv":
        with open(output, "w", newline="", encoding="utf-8-sig") as file:
            writer = csv.DictWriter(file, fieldnames=FIELDS)
            writer.writeheader()
            writer.writerows(articles)
    elif output.suffix.lower() == ".json":
        with open(output, "w", encoding="utf-8") as file:
            json.dump(articles, file, indent=2, ensure_ascii=False)
    elif output.suffix.lower() == ".xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "News Articles"
        sheet.append(FIELDS)

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid", fgColor="1F4E78"
            )

        for article in articles:
            sheet.append([article.get(field, "") or "" for field in FIELDS])

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(output)
    else:
        print("Choose .csv, .xlsx, or .json")
        return

    print(f"Exported {len(articles)} articles to {output.name}")


def main():
    parser = argparse.ArgumentParser(
        description="News Aggregator CLI"
    )
    commands = parser.add_subparsers(dest="command", required=True)

    commands.add_parser("fetch", help="Fetch headlines from RSS feeds")

    list_parser = commands.add_parser("list", help="List saved articles")
    add_filters(list_parser)

    export_parser = commands.add_parser("export", help="Export articles")
    add_filters(export_parser)
    export_parser.add_argument(
        "--output", default="news_export.csv"
    )

    args = parser.parse_args()

    try:
        if args.command == "fetch":
            fetch_news()
        elif args.command == "list":
            list_news(args)
        elif args.command == "export":
            export_news(args)
    except (ValueError, sqlite3.Error) as error:
        parser.error(str(error))


if __name__ == "__main__":
    main()
